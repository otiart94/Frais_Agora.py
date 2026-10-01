import base64
import datetime
import io
import os
import re
import shutil
import zipfile

import openpyxl
import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image
from supabase import Client, create_client

# --- SUPPORT ROBUSTE DES PHOTOS HEIC ---
try:
  from pillow_heif import register_heif_opener

  register_heif_opener()
except ImportError:
  pass

st.set_page_config(
    page_title="Note de Frais Agora - Gestion des Notes de Frais",
    page_icon="💶",
    layout="wide",
)

# --- CONNEXION SUPABASE ---
try:
  SUPABASE_URL = st.secrets["supabase"]["url"]
  SUPABASE_KEY = st.secrets["supabase"]["key"]
  supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)
except Exception as e:
  st.error(
      "Erreur de configuration des secrets Supabase. Vérifie ton fichier"
      f" secrets.toml : {e}"
  )
  st.stop()

# --- CHARTE GRAPHIQUE PROFESSIONNELLE & DESIGN DES CHAMPS & DASHBOARD ---
st.markdown(
    """
    <style>
        [data-testid="stSidebar"] {
            background-color: #18302c !important;
            padding-top: 1rem;
        }
        [data-testid="stSidebar"] *, 
        [data-testid="stSidebar"] label, 
        [data-testid="stSidebar"] span, 
        [data-testid="stSidebar"] p, 
        [data-testid="stSidebar"] div,
        [data-testid="stSidebar"] .stMarkdown {
            color: #ffffff !important;
        }
        [data-testid="stSidebar"] .stButton button {
            background-color: #229954 !important;
            color: #ffffff !important;
            border: none !important;
            border-radius: 6px !important;
            font-weight: 600 !important;
            width: 100%;
        }
        [data-testid="stSidebar"] .stButton button:hover {
            background-color: #1e8449 !important;
            color: #ffffff !important;
        }
        [data-testid="stSidebar"] .stSelectbox div[data-baseweb="select"] {
            background-color: #122421 !important;
            color: #ffffff !important;
            border-color: #2b4d47 !important;
        }
        [data-testid="stSidebar"] .stSelectbox svg {
            fill: #ffffff !important;
        }
        div[data-testid="stForm"] input, 
        div[data-testid="stForm"] div[data-baseweb="select"] div,
        div[data-testid="stForm"] .stDateInput input {
            background-color: #ffffff !important;
            color: #000000 !important;
            border: 1px solid #cbd5e1 !important;
            border-radius: 6px !important;
            box-shadow: 0 1px 2px rgba(0, 0, 0, 0.05) !important;
        }
        div[data-testid="stForm"] input, 
        div[data-testid="stForm"] .stDateInput input {
            color: #000000 !important;
        }
        button[kind="primary"] {
            background-color: #229954 !important;
            border-color: #229954 !important;
            color: #ffffff !important;
        }
        button[kind="primary"]:hover {
            background-color: #1e8449 !important;
            border-color: #1e8449 !important;
        }
        @media (min-width: 769px) {
            div[data-testid="stForm"] input, 
            div[data-testid="stForm"] div[data-baseweb="select"],
            div[data-testid="stForm"] .stDateInput input,
            div[data-testid="stFileUploader"] {
                max-width: 380px !important;
            }
        }
    </style>
    """,
    unsafe_allow_html=True,
)


def _mdp(nom: str) -> str:
  try:
    return st.secrets["passwords"][nom]
  except Exception:
    return "123"


UTILISATEURS = {
    "Sabrina": {"nom_complet": "Sabrina", "password": _mdp("Sabrina")},
    "Alain": {"nom_complet": "Alain Autier", "password": _mdp("Alain")},
    "Edouard": {"nom_complet": "Edouard Dupont", "password": _mdp("Alain")},
    
}

COLS_MONTANTS = [
    "parking",
    "entretien_materiel",
    "reception",
    "hotel",
    "restauration",
    "gasoil",
    "tva",
    "total",
]


# --- VISUALISEUR INTERACTIF AVEC ZOOM FLUIDE ---
def afficher_image_zoomable(img_b64: str, hauteur: int = 500):
  html_code = (
      """
    <div id="v" style="width:100%; height:__H__px; overflow:hidden; position:relative; touch-action:none; border:1px solid #ddd; border-radius:8px; background:#111; display:flex; align-items:center; justify-content:center;">
      <img id="im" draggable="false" src="data:image/jpeg;base64,__IMG__" style="max-width:100%; max-height:100%; object-fit:contain; cursor:grab; user-select:none; -webkit-user-drag:none; transform-origin:center center; transition: transform 0.05s ease-out;">
    </div>
    <script>
      const v = document.getElementById('v'), im = document.getElementById('im');
      let s = 1, x = 0, y = 0;
      const pts = new Map();
      function apply() { im.style.transform = 'translate(' + x + 'px, ' + y + 'px) scale(' + s + ')'; }
      function zoomAt(cx, cy, factor) {
        const s2 = Math.min(8, Math.max(1, s * factor));
        x = cx - (cx - x) * (s2 / s);
        y = cy - (cy - y) * (s2 / s);
        s = s2;
      }
      function info() {
        const a = [...pts.values()];
        const r = v.getBoundingClientRect();
        return { d: Math.hypot(a[0].x - a[1].x, a[0].y - a[1].y), mx: (a[0].x + a[1].x) / 2 - r.left, my: (a[0].y + a[1].y) / 2 - r.top };
      }
      v.addEventListener('pointerdown', e => { v.setPointerCapture(e.pointerId); pts.set(e.pointerId, {x: e.clientX, y: e.clientY}); im.style.cursor='grabbing'; });
      v.addEventListener('pointermove', e => {
        if (!pts.has(e.pointerId)) return;
        const prev = pts.get(e.pointerId);
        if (pts.size === 2) {
          const before = info();
          pts.set(e.pointerId, {x: e.clientX, y: e.clientY});
          const after = info();
          zoomAt(after.mx, after.my, after.d / before.d);
          x += after.mx - before.mx; y += after.my - before.my;
        } else {
          pts.set(e.pointerId, {x: e.clientX, y: e.clientY});
          if (s > 1) { x += e.clientX - prev.x; y += e.clientY - prev.y; }
        }
        apply();
      });
      ['pointerup', 'pointercancel', 'pointerleave'].forEach(ev => v.addEventListener(ev, e => { pts.delete(e.pointerId); im.style.cursor='grab'; }));
      v.addEventListener('wheel', e => {
        e.preventDefault();
        const r = v.getBoundingClientRect();
        const cx = e.clientX - r.left, cy = e.clientY - r.top;
        const factor = e.deltaY < 0 ? 1.15 : 0.85;
        zoomAt(cx, cy, factor);
        apply();
      }, { passive: false });
    </script>
    """
      .replace("__IMG__", img_b64)
      .replace("__H__", str(hauteur))
  )
  components.html(html_code, height=hauteur + 20)


# --- AUTHENTIFICATION ---
if "authentifie" not in st.session_state:
  st.session_state["authentifie"] = False
  st.session_state["user_key"] = ""
  st.session_state["user_nom_complet"] = ""

if not st.session_state["authentifie"]:
  st.title("🔐 Connexion à Frais Agora")
  with st.form("form_login"):
    username_input = st.text_input("Identifiant (Sabrina ou Alain)")
    password_input = st.text_input("Mot de passe", type="password")
    if st.form_submit_button("Se connecter"):
      user_trouve = next(
          (
              u
              for u in UTILISATEURS
              if u.lower() == username_input.strip().lower()
          ),
          None,
      )
      if user_trouve and UTILISATEURS[user_trouve]["password"] == password_input:
        st.session_state["authentifie"] = True
        st.session_state["user_key"] = user_trouve.lower()
        st.session_state["user_nom_complet"] = UTILISATEURS[user_trouve][
            "nom_complet"
        ]
        st.rerun()
      else:
        st.error("Identifiant ou mot de passe incorrect.")
  st.stop()

user_key = st.session_state["user_key"]
nom_salarie_defaut = st.session_state["user_nom_complet"]

colonnes_attendues = [
    "id",
    "user_key",
    "periode",
    "nom",
    "date",
    "libelle",
    "chantier",
    "parking",
    "entretien_materiel",
    "reception",
    "hotel",
    "restauration",
    "gasoil",
    "tva",
    "total",
    "justificatif",
]

# Chargement initial des données depuis Supabase
if "frais_data" not in st.session_state:
  try:
    response = (
        supabase.table("frais_depenses")
        .select("*")
        .eq("user_key", user_key)
        .execute()
    )
    if response.data:
      st.session_state["frais_data"] = pd.DataFrame(response.data)
      st.session_state["id_counter"] = (
          int(st.session_state["frais_data"]["id"].max())
          if "id" in st.session_state["frais_data"].columns
          and not st.session_state["frais_data"].empty
          else len(st.session_state["frais_data"])
      )
    else:
      st.session_state["frais_data"] = pd.DataFrame(columns=colonnes_attendues)
      st.session_state["id_counter"] = 0
  except Exception as e:
    st.error(f"Erreur de chargement depuis Supabase : {e}")
    st.session_state["frais_data"] = pd.DataFrame(columns=colonnes_attendues)
    st.session_state["id_counter"] = 0

# --- BARRE LATÉRALE : Gestion des Archives & Paiement ---
with st.sidebar:
  st.write(f"👤 Connecté en tant que : **{nom_salarie_defaut}**")
  if st.button("🚪 Se déconnecter"):
    for k in ("authentifie", "user_key", "user_nom_complet", "frais_data"):
      st.session_state.pop(k, None)
    st.rerun()
  st.markdown("---")

  st.header("📂 Gestion des Mois")
  try:
    res_arc = (
        supabase.table("frais_depenses")
        .select("periode")
        .eq("user_key", user_key)
        .execute()
    )
    if res_arc.data:
      periodes_uniques = list(
          set([row["periode"] for row in res_arc.data if row.get("periode")])
      )
      if periodes_uniques:
        periode_choisie = st.selectbox(
            "Filtrer par Période",
            ["-- Toutes les dépenses --"] + periodes_uniques,
        )
        if periode_choisie != "-- Toutes les dépenses --":
          if st.button("📂 Appliquer le filtre"):
            res_filtre = (
                supabase.table("frais_depenses")
                .select("*")
                .eq("user_key", user_key)
                .eq("periode", periode_choisie)
                .execute()
            )
            st.session_state["frais_data"] = pd.DataFrame(res_filtre.data)
            st.rerun()
  except Exception:
    pass

  # --- SECTION COMMERCIALE / STRIPE ---
  st.markdown("---")
  st.markdown("### 💎 Obtenir l'accès Pro")
  st.markdown(
      "Profitez de Frais Agora en illimité pour gérer vos notes de frais :"
  )
  st.markdown(
      "[👉 S'abonner (5,00 €/mois)]"
      "(https://buy.stripe.com/test_dRmfZi3KO2PR2Yx9wfaAw00)"
  )

st.markdown("---")
st.title("📄 Note de Frais - Suivi des Frais (Espace Pro)")
st.markdown("---")

if "uploader_key" not in st.session_state:
  st.session_state["uploader_key"] = 0

st.subheader("📸 1. Joignez ou scannez votre ticket de caisse")
col_upl1, col_upl2 = st.columns([1.1, 1.9])
with col_upl1:
  photo_uploadee = st.file_uploader(
      "Ticket de caisse",
      type=["png", "jpg", "jpeg", "heic", "HEIC"],
      key=f"uploader_ticket_{st.session_state['uploader_key']}",
  )
with col_upl2:
  st.write("")
  st.write("")
  if st.button("🔄 Nouveau ticket", type="secondary"):
    st.session_state.pop("ticket_actuel_bytes", None)
    st.session_state["uploader_key"] += 1
    st.rerun()

# Stockage contrôlé des octets du fichier en session
if photo_uploadee is not None:
  st.session_state["ticket_actuel_bytes"] = photo_uploadee.getvalue()

image_rgb = None
col_form, col_ticket = st.columns([1.2, 1], gap="large")

with col_ticket:
  st.subheader("👁️ Visualisation (Zoom interactif)")
  angle_rotation = st.selectbox(
      "🔄 Pivoter", [0, 90, 180, 270], index=0, key="select_rotation"
  )
  if "ticket_actuel_bytes" in st.session_state and st.session_state[
      "ticket_actuel_bytes"
  ]:
    try:
      image_pil = Image.open(
          io.BytesIO(st.session_state["ticket_actuel_bytes"])
      )
      if angle_rotation > 0:
        image_pil = image_pil.rotate(-angle_rotation, expand=True)
      image_rgb = image_pil.convert("RGB")
    except Exception as e:
      st.error(f"Erreur image : {e}")

  if image_rgb is not None:
    buffered = io.BytesIO()
    image_rgb.save(buffered, format="JPEG", quality=95)
    afficher_image_zoomable(base64.b64encode(buffered.getvalue()).decode())

with col_form:
  st.subheader("✍ 2. Saisie de la dépense")
  with st.form("form_frais", clear_on_submit=True):
    col_f1, col_f2 = st.columns(2)
    with col_f1:
      periode = st.text_input("Période (ex: août-26)", value="août-26")
      nom = st.text_input("Nom du salarié", value=nom_salarie_defaut)
      date_frais = st.date_input("Date de la dépense", value=datetime.date.today())
      libelle = st.text_input("Libellé / Type", value="")
    with col_f2:
      categorie = st.selectbox(
          "Catégorie",
          [
              "Parking / Péage",
              "Restauration",
              "Gasoil",
              "Hôtel",
              "Entretien matériel",
              "Réception",
          ],
      )
      chantier = st.text_input("Chantier (ex: AEU, Inspire)", value="")
      montant_str = st.text_input("Montant TTC (€)", value="0.00")
      tva_str = st.text_input("TVA (€)", value="0.00")

    submitted = st.form_submit_button("💾 Enregistrer définitivement la dépense")

    if submitted:
      try:
        montant = float(str(montant_str).replace(",", "."))
        tva = float(str(tva_str).replace(",", "."))
      except ValueError:
        montant, tva = 0.0, 0.0

      if montant > 0:
        st.session_state["id_counter"] += 1

        nom_fichier_sauvegarde = "Aucun justificatif"
        if (
            "ticket_actuel_bytes" in st.session_state
            and st.session_state["ticket_actuel_bytes"]
        ):
          try:
            timestamp_str = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
            nom_fichier_sauvegarde = f"ticket_{user_key}_{timestamp_str}.jpg"

            img_to_upload = Image.open(
                io.BytesIO(st.session_state["ticket_actuel_bytes"])
            )
            if angle_rotation > 0:
              img_to_upload = img_to_upload.rotate(-angle_rotation, expand=True)
            buf_final = io.BytesIO()
            img_to_upload.convert("RGB").save(
                buf_final, format="JPEG", quality=95
            )

            supabase.storage.from_("justificatifs").upload(
                path=nom_fichier_sauvegarde,
                file=buf_final.getvalue(),
                file_options={"content-type": "image/jpeg", "upsert": "true"},
            )
          except Exception as img_err:
            st.error(f"❌ ERREUR UPLOAD STORAGE : {img_err}")
            st.stop()

        nouvelle_ligne = {
            "id": int(st.session_state["id_counter"]),
            "user_key": str(user_key),
            "periode": str(periode),
            "nom": str(nom),
            "date": str(date_frais.strftime("%d-%b")),
            "libelle": str(libelle if libelle else categorie),
            "chantier": str(chantier),
            "parking": float(montant if categorie == "Parking / Péage" else 0.0),
            "entretien_materiel": float(
                montant if categorie == "Entretien matériel" else 0.0
            ),
            "reception": float(montant if categorie == "Réception" else 0.0),
            "hotel": float(montant if categorie == "Hôtel" else 0.0),
            "restauration": float(
                montant if categorie == "Restauration" else 0.0
            ),
            "gasoil": float(montant if categorie == "Gasoil" else 0.0),
            "tva": float(max(tva, 0.0)),
            "total": float(montant),
            "justificatif": str(nom_fichier_sauvegarde),
        }

        try:
          supabase.table("frais_depenses").insert(nouvelle_ligne).execute()
          st.success("✅ Dépense et justificatif enregistrés dans le Cloud !")
          st.session_state.pop("ticket_actuel_bytes", None)

          response = (
              supabase.table("frais_depenses")
              .select("*")
              .eq("user_key", user_key)
              .execute()
          )
          st.session_state["frais_data"] = pd.DataFrame(response.data)
          st.rerun()
        except Exception as db_err:
          st.error(f"❌ ERREUR EXACTE CAPTURÉE : {db_err}")
      else:
        st.error("Veuillez indiquer un montant supérieur à 0.")

# --- TABLEAU RÉCAPITULATIF & MINI-TABLEAU DE BORD (KPIs) ---
st.markdown("---")
nb_lignes = len(st.session_state["frais_data"])
texte_compteur = (
    f" ({nb_lignes} dépense{'s' if nb_lignes > 1 else ''} enregistrée{'s' if nb_lignes > 1 else ''})"
    if nb_lignes > 0
    else " (Aucune dépense)"
)
st.subheader(
    f"📊 Tableau récapitulatif des frais de : {nom_salarie_defaut}"
    f" {texte_compteur}"
)

if not st.session_state["frais_data"].empty:
  df_actuel = st.session_state["frais_data"]
  tot_parking = (
      df_actuel["parking"].sum() if "parking" in df_actuel.columns else 0.0
  )
  tot_restauration = (
      df_actuel["restauration"].sum()
      if "restauration" in df_actuel.columns
      else 0.0
  )
  tot_gasoil = (
      df_actuel["gasoil"].sum() if "gasoil" in df_actuel.columns else 0.0
  )
  tot_hotel = df_actuel["hotel"].sum() if "hotel" in df_actuel.columns else 0.0
  tot_entretien = (
      df_actuel["entretien_materiel"].sum()
      if "entretien_materiel" in df_actuel.columns
      else 0.0
  )
  tot_reception = (
      df_actuel["reception"].sum() if "reception" in df_actuel.columns else 0.0
  )
  total_general = df_actuel["total"].sum() if "total" in df_actuel.columns else 0.0

  kpi_col1, kpi_col2, kpi_col3, kpi_col4 = st.columns(4)
  with kpi_col1:
    st.metric(label="💶 Total Général", value=f"{total_general:.2f} €")
  with kpi_col2:
    st.metric(label="🍽️️ Restauration", value=f"{tot_restauration:.2f} €")
  with kpi_col3:
    st.metric(label="⛽ Gasoil", value=f"{tot_gasoil:.2f} €")
  with kpi_col4:
    st.metric(label="🅿️ Parking / Péage", value=f"{tot_parking:.2f} €")

  with st.expander("Voir le détail des autres catégories (Hôtel, Entretien, Réception)"):
    kpi_sub1, kpi_sub2, kpi_sub3 = st.columns(3)
    with kpi_sub1:
      st.metric(label="🏨 Hôtel", value=f"{tot_hotel:.2f} €")
    with kpi_sub2:
      st.metric(label="🔧 Entretien matériel", value=f"{tot_entretien:.2f} €")
    with kpi_sub3:
      st.metric(label="🥂 Réception", value=f"{tot_reception:.2f} €")

  st.markdown("---")
  df_travail = st.session_state["frais_data"].copy()
  if "❌ Suppr" not in df_travail.columns:
    df_travail.insert(0, "❌ Suppr", False)

  df_affiche = (
      df_travail.drop(columns=["id"]) if "id" in df_travail.columns else df_travail
  )
  df_modifie = st.data_editor(df_affiche, use_container_width=True)

  # --- ACTIONS BAS DE PAGE : SUPPRESSION & DÉCLARATION OFFICIELLE ---
  st.markdown("---")
  col_bas1, col_bas2 = st.columns(2)

  with col_bas1:
    st.subheader("🗑 Actions sur les lignes")
    if st.button("🗑️ Supprimer les lignes cochées dans le tableau"):
      lignes_a_cocher = df_modifie["❌ Suppr"] == True
      if lignes_a_cocher.any():
        lignes_a_supprimer_df = df_travail.loc[lignes_a_cocher]
        for _, row_del in lignes_a_supprimer_df.iterrows():
          row_id = row_del.get("id")
          nom_fic = row_del.get("justificatif")
          if pd.notna(row_id):
            try:
              supabase.table("frais_depenses").delete().eq(
                  "id", int(row_id)
              ).eq("user_key", user_key).execute()
            except Exception:
              pass
          if pd.notna(nom_fic) and nom_fic != "Aucun justificatif":
            try:
              supabase.storage.from_("justificatifs").remove([str(nom_fic)])
            except Exception:
              pass

        response = (
            supabase.table("frais_depenses")
            .select("*")
            .eq("user_key", user_key)
            .execute()
        )
        st.session_state["frais_data"] = pd.DataFrame(response.data)
        st.success("Ligne(s) et justificatif(s) supprimés avec succès !")
        st.rerun()
      else:
        st.warning("Veuillez cocher au moins une case ❌ Suppr.")

  with col_bas2:
    st.subheader("📁 Déclaration officielle")
    if st.button("📄 Préparer les fichiers de la note de frais", type="primary"):
      wb = openpyxl.Workbook()
      ws = wb.active
      ws.title = "Note de Frais"

      font_titre = Font(name="Arial", size=14, bold=True, color="1F4E78")
      font_sous_titre = Font(name="Arial", size=11, bold=True, color="333333")
      font_en_tete = Font(name="Arial", size=10, bold=True, color="FFFFFF")
      fill_en_tete = PatternFill(
          start_color="1F4E78", end_color="1F4E78", fill_type="solid"
      )
      font_donnees = Font(name="Arial", size=10)
      font_total = Font(name="Arial", size=10, bold=True)
      align_centre = Alignment(horizontal="center", vertical="center")
      align_droite = Alignment(horizontal="right", vertical="center")
      align_gauche = Alignment(horizontal="left", vertical="center")
      gris = Side(style="thin", color="D9D9D9")
      bordure_fine = Border(left=gris, right=gris, top=gris, bottom=gris)
      bordure_total = Border(
          top=Side(style="thin", color="000000"),
          bottom=Side(style="double", color="000000"),
      )

      periode_rapport = (
          str(df_modifie["periode"].iloc[0])
          if "periode" in df_modifie.columns and not df_modifie.empty
          else "Mois"
      )

      ws["A1"] = f"NOTE DE FRAIS - {nom_salarie_defaut.upper()}"
      ws["A1"].font = font_titre
      ws["A2"] = f"Salarié(e) : {nom_salarie_defaut} | Période : {periode_rapport}"
      ws["A2"].font = font_sous_titre
      ws.append([])

      colonnes_excel = [c for c in colonnes_attendues if c != "id"]
      ws.append(colonnes_excel)
      for col_num in range(1, len(colonnes_excel) + 1):
        cell = ws.cell(row=4, column=col_num)
        cell.font = font_en_tete
        cell.fill = fill_en_tete
        cell.alignment = align_centre
        cell.border = bordure_fine

      for _, row in st.session_state["frais_data"].iterrows():
        valeurs = []
        for c_name in colonnes_excel:
          val = row.get(c_name, "")
          if c_name in COLS_MONTANTS:
            try:
              val = float(val) if pd.notna(val) and val != "" else 0.0
            except ValueError:
              val = 0.0
          elif pd.isna(val):
            val = ""
          valeurs.append(val)
        ws.append(valeurs)

        r = ws.max_row
        for col_num, c_name in enumerate(colonnes_excel, start=1):
          cell = ws.cell(row=r, column=col_num)
          cell.font = font_donnees
          cell.border = bordure_fine
          if c_name in COLS_MONTANTS:
            cell.number_format = "#,##0.00 €"
            cell.alignment = align_droite
          else:
            cell.alignment = align_gauche

      ligne_totaux = []
      for c_name in colonnes_excel:
        if c_name in COLS_MONTANTS:
          ligne_totaux.append(
              pd.to_numeric(
                  st.session_state["frais_data"][c_name], errors="coerce"
              ).sum()
          )
        elif c_name == "libelle":
          ligne_totaux.append("TOTAL GÉNÉRAL")
        else:
          ligne_totaux.append("")
      ws.append(ligne_totaux)

      tot_row = ws.max_row
      for col_num, c_name in enumerate(colonnes_excel, start=1):
        cell = ws.cell(row=tot_row, column=col_num)
        cell.font = font_total
        cell.border = bordure_total
        if c_name in COLS_MONTANTS:
          cell.number_format = "#,##0.00 €"
          cell.alignment = align_droite
        else:
          cell.alignment = align_gauche

      for col in ws.columns:
        max_len = max(
            (len(str(c.value)) for c in col if c.value is not None), default=0
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(
            max_len + 4, 12
        )

      output_excel = io.BytesIO()
      wb.save(output_excel)
      st.session_state["excel_bytes"] = output_excel.getvalue()
      st.session_state["excel_nom"] = (
          f"NDF-{nom_salarie_defaut.replace(' ', '-')}-{periode_rapport}.xlsx"
      )

      memory_zip = io.BytesIO()
      with zipfile.ZipFile(memory_zip, "w", zipfile.ZIP_DEFLATED) as zf:
        fichiers_actifs = st.session_state["frais_data"][
            "justificatif"
        ].tolist()
        for nom_fic in fichiers_actifs:
          if pd.notna(nom_fic) and nom_fic != "Aucun justificatif":
            try:
              res_file = supabase.storage.from_("justificatifs").download(
                  str(nom_fic)
              )
              if res_file:
                if isinstance(res_file, bytes):
                  zf.writestr(str(nom_fic), res_file)
                else:
                  zf.writestr(str(nom_fic), bytes(res_file))
            except Exception:
              pass
      memory_zip.seek(0)
      st.session_state["zip_bytes"] = memory_zip.getvalue()
      st.session_state["zip_nom"] = (
          f"Justificatifs-{nom_salarie_defaut.replace(' ', '-')}-{periode_rapport}.zip"
      )

    if st.session_state.get("excel_bytes"):
      st.success("🎉 Fichiers de déclaration générés avec succès !")
      col_dl1, col_dl2 = st.columns(2)
      with col_dl1:
        st.download_button(
            label="📥 Télécharger l'Excel (.xlsx)",
            data=st.session_state["excel_bytes"],
            file_name=st.session_state.get(
                "excel_nom", "note_de_frais.xlsx"
            ),
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
      with col_dl2:
        if st.session_state.get("zip_bytes"):
          st.download_button(
              label="📥 Télécharger tous les Tickets (.zip)",
              data=st.session_state["zip_bytes"],
              file_name=st.session_state.get(
                  "zip_nom", "justificatifs.zip"
              ),
              mime="application/zip",
          )
else:
  st.info("Aucune dépense enregistrée pour le moment.")
